from pathlib import Path
import zipfile

import pytest

from app.indexing import DocumentExtractor
from app.lifecycle import enqueue_reindex, enqueue_retry
from app.models import Document, DocumentVersion


def test_extractor_rejects_excessive_text(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("x" * 101, encoding="utf-8")

    with pytest.raises(ValueError, match="extracted-content limit"):
        DocumentExtractor(max_tokens=25).extract(path, "text/plain")


def test_reindex_and_retry_limits_are_enforced():
    document = Document(original_filename="notes.txt")
    document.versions.extend(DocumentVersion(version_number=i) for i in range(1, 3))
    with pytest.raises(ValueError, match="maximum number"):
        enqueue_reindex(document, max_versions=2)

    version = DocumentVersion(version_number=1, status="failed")
    version.ingestion_jobs.extend([])
    enqueue_retry(version, max_jobs=1)
    with pytest.raises(ValueError, match="maximum number"):
        enqueue_retry(version, max_jobs=1)
def test_docx_archive_budget_is_checked_before_parser(tmp_path: Path):
    path = tmp_path / "oversized.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "x" * 101)

    with pytest.raises(ValueError, match="extracted-content limit"):
        DocumentExtractor(max_tokens=25, max_archive_bytes=100).extract(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

def test_normal_docx_remains_supported(tmp_path: Path):
    from docx import Document as WordDocument

    path = tmp_path / "notes.docx"
    document = WordDocument()
    document.add_paragraph("A short supported document.")
    document.save(path)

    sections = DocumentExtractor(max_tokens=25).extract(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    assert [section.text for section in sections] == ["A short supported document."]
